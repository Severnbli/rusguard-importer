# -*- coding: utf-8 -*-
"""Строки таблицы приложения, их проверка, а также необязательные шаблоны/импорт из Excel."""

from __future__ import annotations

import os
import re
import shutil
import subprocess
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from keys import parse_keys

# в собранном PyInstaller-приложении ресурсы распакованы в sys._MEIPASS
BASE_DIR = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parents[1]))
TEMPLATES_DIR = BASE_DIR / "templates"
MODE_ADD = "add"
MODE_NEW = "new"
MODE_ACCESS = "access"  # только просмотр и выгрузка карточек в Access
TEMPLATES = {MODE_ADD: TEMPLATES_DIR / "template_add.xlsx", MODE_NEW: TEMPLATES_DIR / "template_new.xlsx"}
ACCESS_TEMPLATE = TEMPLATES_DIR / "template.accdb"

_MAX_ROWS = 1000
_LISTS_SHEET = "Справочники"


@dataclass
class Entry:
    """Строка таблицы в окне приложения (то, что вводит пользователь)."""
    fullname: str = ""
    keys: str = ""
    position: str = ""
    group: str = ""
    photo: Optional[bytes] = None
    employee_id: Optional[str] = None  # задан, если сотрудник выбран из БД
    # исходные данные профиля из БД (режим редактирования) — чтобы отправлять только изменения
    photo_original: Optional[bytes] = None
    position_original: str = ""
    group_original: str = ""
    current_keys: list = field(default_factory=list)  # ключи сотрудника в БД (числа)
    fullname_original: str = ""


@dataclass
class Row:
    """Проверенная строка, готовая к импорту."""
    number: int
    fullname: str
    lastname: str = ""
    firstname: str = ""
    secondname: Optional[str] = None
    keys: list = field(default_factory=list)
    position_id: Optional[str] = None
    group_id: Optional[str] = None
    photo: Optional[bytes] = None
    employee_id: Optional[str] = None
    remove_keys: list = field(default_factory=list)
    rename: bool = False
    errors: list = field(default_factory=list)
    warnings: list = field(default_factory=list)

    @property
    def label(self) -> str:
        return f"строка {self.number} ({self.fullname or 'без ФИО'})"


def split_fio(fullname: str):
    """-> (фамилия, имя, отчество|None, предупреждения)."""
    tokens = [t for t in re.sub(r"\s+", " ", fullname.strip()).split(" ") if t]
    warnings = []
    if len(tokens) >= 3:
        if len(tokens) > 3:
            warnings.append(f"в ФИО {len(tokens)} слов — отчество записано как '{' '.join(tokens[2:])}'")
        return tokens[0], tokens[1], " ".join(tokens[2:]), warnings
    if len(tokens) == 2:
        warnings.append("отчество не указано")
        return tokens[0], tokens[1], None, warnings
    warnings.append("в ФИО только одно слово — записано как фамилия")
    return (tokens[0] if tokens else ""), "", None, warnings


def normalize_name(text: str) -> str:
    return re.sub(r"\s+", " ", str(text).strip()).lower()


# ----------------------------------------------------------------------
#  Проверка строк таблицы
# ----------------------------------------------------------------------

def _name_lookup(items: list) -> dict:
    result: dict = {}
    for item_id, name in items:
        if name:
            result.setdefault(normalize_name(name), set()).add(str(item_id))
    return result


def build_rows(entries: list, mode: str, key_format: str, positions: list, groups: list) -> list:
    pos_lookup, grp_lookup = _name_lookup(positions), _name_lookup(groups)
    rows, seen = [], {}

    for number, entry in enumerate(entries, start=1):
        fullname = entry.fullname.strip()
        if not fullname and not entry.keys.strip():
            continue
        if fullname.startswith("#"):
            continue  # закомментированная строка

        row = Row(number=number, fullname=fullname, photo=entry.photo, employee_id=entry.employee_id)
        rows.append(row)
        if not fullname:
            row.errors.append("не заполнено ФИО")
            continue
        row.lastname, row.firstname, row.secondname, warns = split_fio(fullname)
        if not entry.employee_id:
            row.warnings.extend(warns)

        row.keys, key_errors = parse_keys(entry.keys, key_format)
        row.errors.extend(f"ключ: {e}" for e in key_errors)

        if mode == MODE_ADD:
            row.photo = entry.photo if entry.photo and entry.photo != entry.photo_original else None
            if entry.employee_id:
                # для сотрудника из БД поле ключей — полный список: разница с БД = добавить/снять
                wanted = row.keys
                row.keys = [k for k in wanted if k not in entry.current_keys]
                if not key_errors:
                    row.remove_keys = [k for k in entry.current_keys if k not in wanted]
                row.rename = fullname != entry.fullname_original.strip()
            if entry.position.strip() != entry.position_original.strip():
                row.position_id = _resolve(entry.position, pos_lookup, "должность", row)
            if entry.group.strip() != entry.group_original.strip():
                row.group_id = _resolve(entry.group, grp_lookup, "группа", row)
            if not (row.keys or row.remove_keys or row.rename or row.photo or row.position_id or row.group_id
                    or row.errors):
                row.errors.append("нет изменений")
            continue

        norm = normalize_name(fullname)
        if norm in seen:
            row.errors.append(f"ФИО повторяется (первый раз — строка {seen[norm]})")
        else:
            seen[norm] = number
        row.position_id = _resolve(entry.position, pos_lookup, "должность", row)
        row.group_id = _resolve(entry.group, grp_lookup, "группа", row)
    return rows


def _resolve(raw, lookup: dict, label: str, row: Row) -> Optional[str]:
    name = str(raw).strip()
    if not name:
        row.errors.append(f"не заполнена {label}")
        return None
    ids = lookup.get(normalize_name(name))
    if not ids:
        row.errors.append(f"{label} '{name}' не найдена на сервере")
        return None
    if len(ids) > 1:
        row.errors.append(f"на сервере несколько записей {label} с названием '{name}'")
        return None
    return next(iter(ids))


# ----------------------------------------------------------------------
#  Excel: шаблон и загрузка строк в таблицу (необязательно)
# ----------------------------------------------------------------------

def _header_map(ws) -> dict:
    return {str(c.value).strip().lower(): c.column for c in ws[1] if c.value}


def create_from_template(mode: str, dest: str, positions: list, groups: list) -> None:
    """Копия шаблона; для режима «новый профиль» Должность/Группа — выпадающие списки с данными сервера."""
    src = TEMPLATES[mode]
    if not src.exists():
        raise FileNotFoundError(f"Не найден шаблон: {src}")
    shutil.copyfile(src, dest)
    wb = load_workbook(dest)
    ws = wb.worksheets[0]

    header = _header_map(ws)
    key_col = header.get("ключ")
    if key_col:  # текст, чтобы Excel не съедал ведущие нули HEX-номеров
        for r in range(2, _MAX_ROWS + 1):
            ws.cell(r, key_col).number_format = "@"

    if mode == MODE_NEW:
        lists = wb.create_sheet(_LISTS_SHEET)
        lists["A1"], lists["B1"] = "Должности", "Группы"
        pos_names = sorted({n for _, n in positions if n})
        grp_names = sorted({n for _, n in groups if n})
        for i, name in enumerate(pos_names, start=2):
            lists.cell(i, 1, name)
        for i, name in enumerate(grp_names, start=2):
            lists.cell(i, 2, name)
        for title, letter, count in (("должность", "A", len(pos_names)), ("группа", "B", len(grp_names))):
            col = header.get(title)
            if not col or not count:
                continue
            dv = DataValidation(
                type="list", formula1=f"='{_LISTS_SHEET}'!${letter}$2:${letter}${count + 1}",
                allow_blank=True, showErrorMessage=True,
            )
            ws.add_data_validation(dv)
            col_letter = ws.cell(1, col).column_letter
            dv.add(f"{col_letter}2:{col_letter}{_MAX_ROWS}")
        lists.sheet_state = "hidden"
    wb.save(dest)


def open_in_editor(path: str) -> bool:
    try:
        if os.name == "nt":
            os.startfile(path)  # type: ignore[attr-defined]
        else:
            subprocess.Popen(["open" if sys.platform == "darwin" else "xdg-open", path])
    except OSError:
        return False
    return True


def load_xlsx_entries(path: str) -> list:
    """Читает столбцы ФИО, Ключ, Должность, Группа первого листа в строки таблицы."""
    ws = load_workbook(path, data_only=True).worksheets[0]
    header = _header_map(ws)
    missing = [h for h in ("фио", "ключ") if h not in header]
    if missing:
        raise ValueError(f"В файле нет столбцов: {', '.join(missing)}.")

    entries = []
    for cells in ws.iter_rows(min_row=2):
        def val(title) -> str:
            v = cells[header[title] - 1].value if title in header else None
            if v is None:
                return ""
            if isinstance(v, float) and v.is_integer():
                v = int(v)
            return str(v).strip()

        entry = Entry(fullname=val("фио"), keys=val("ключ"), position=val("должность"), group=val("группа"))
        if entry.fullname or entry.keys:
            entries.append(entry)
    return entries
